# -*- coding: utf-8 -*-
"""国考职位筛选：硬条件判定、三档分级、xlsx 与复核清单输出。"""

import argparse
import glob
import json
import os
import re
import sys
from collections import Counter, OrderedDict
from datetime import datetime

from jobs import (cell_text, code_text, load_jobs, normalize, parse_degree,
                  parse_education, parse_years, select_major_segment,
                  split_sentences)

TIER_ORDER = {"优先推荐": 0, "待复核": 1, "不推荐": 2}
TIER_FILL = {"优先推荐": "C6EFCE", "待复核": "FFEB9C", "不推荐": "FFC7CE"}
TIER_FONT = {"优先推荐": "006100", "待复核": "9C6500", "不推荐": "9C0006"}

CONDITION_DEFAULTS = {
    "education": "",
    "degree": "",
    "major": [],
    "major_codes": [],
    "political_status": "群众",
    "grassroots_years": 0,
    "service_program": "无",
    "fresh_graduate": False,
    "graduate_year": None,
    "regions": [],
    "english": {},
    "accept_service_years": True,
    "max_service_years": None,
    "accept_intensity": True,
    "gender": "",
    "hukou": "",
}

SERVICE_KEYWORDS = [
    ("大学生村官", ("大学生村官", "村官")),
    ("三支一扶", ("三支一扶",)),
    ("西部计划", ("西部计划", "志愿服务西部")),
    ("特岗教师", ("特岗教师",)),
    ("退役士兵", ("退役士兵", "退役大学生士兵")),
]

CERT_PATTERN = (r"(法律职业资格|注册会计师|CPA|税务师|执业医师|执业药师|"
                r"教师资格|计算机等级|专业四级|专业八级|资格证书|职业资格|"
                r"会计专业技术资格|注册建筑师|注册安全工程师)")
PHYSICAL_PATTERN = (r"(体能测评|体能测试|体检|视力|色盲|色弱|身高|体重|"
                    r"人民警察|警察体能)")
POLITICAL_PATTERN = r"(政治考察|政审|政治审查)"
HUKOU_PATTERN = r"(户籍|生源|户口)"
GENDER_PATTERN = r"(男性|女性|适合男性|适合女性|限男性|限女性)"
INTENSITY_PATTERN = (r"(经常出差|长期出差|出差|夜间值班|值夜班|加班|高强度|"
                     r"野外作业|艰苦边远|高原|海岛|轮班|驻外|封闭管理)")
MAJOR_CATEGORY_PATTERN = r"(类|门类|一级学科|相关专业|相近专业|专业目录)"


def issue(category, detail, field, evidence, level="fail"):
    return {
        "category": category,
        "detail": detail,
        "field": field,
        "evidence": cell_text(evidence),
        "level": level,
    }


def evidence_for(text, pattern):
    for sentence in split_sentences(text):
        if re.search(pattern, sentence):
            return sentence
    return ""


def user_edu_level(conditions):
    if isinstance(conditions.get("education_level"), int):
        return conditions["education_level"]
    text = normalize(conditions.get("education", ""))
    if "博士" in text:
        return 4
    if "硕士" in text or "研究生" in text:
        return 3
    if "本科" in text:
        return 2
    if "大专" in text or "专科" in text:
        return 1
    return 0


def user_degree_level(conditions):
    if isinstance(conditions.get("degree_level"), int):
        return conditions["degree_level"]
    text = normalize(conditions.get("degree", ""))
    if "博士" in text:
        return 3
    if "硕士" in text:
        return 2
    if "学士" in text:
        return 1
    if text in ("无", "无学位", "没有"):
        return 0
    # 未填写学位时，按学历推定：本科=学士，硕士=硕士，博士=博士。
    return {2: 1, 3: 2, 4: 3}.get(user_edu_level(conditions), 0)


def check_education(job, conditions):
    raw = job["education_text"]
    parsed = parse_education(raw)
    levels = parsed["levels"]
    if levels is None:
        return []
    if not levels:
        return [issue("学历", "学历要求无法解析，需人工确认", "学历", raw, "review")]
    if user_edu_level(conditions) not in levels:
        return [issue("学历", "考生学历为%s，职位要求%s" % (
            conditions.get("education") or "未填写", raw), "学历", raw)]
    return []


def check_degree(job, conditions):
    raw = job["degree_text"]
    parsed = parse_degree(raw)
    levels = parsed["levels"]
    if parsed["matching"]:
        if user_degree_level(conditions) <= 0:
            return [issue("学位", "职位要求与最高学历相对应的学位，考生未提供学位", "学位", raw)]
        return []
    if levels is None:
        return []
    if not levels:
        return [issue("学位", "学位要求无法解析，需人工确认", "学位", raw, "review")]
    if user_degree_level(conditions) not in levels:
        return [issue("学位", "考生学位不符合职位要求", "学位", raw)]
    return []


def check_major(job, conditions):
    raw = job["major_text"]
    compact = normalize(raw)
    if not compact or compact in ("不限", "无限制", "专业不限", "不限专业"):
        return []

    names = [normalize(name) for name in conditions.get("major") or [] if name]
    codes = [normalize(str(code)) for code in conditions.get("major_codes") or [] if code]

    segment, segmented, matched_level = select_major_segment(raw, user_edu_level(conditions))
    haystack = normalize(segment)

    if not haystack or haystack in ("不限", "无限制", "专业不限", "不限专业"):
        return []

    if any(code in haystack for code in codes):
        return []

    tokens = [token for token in re.split(r"[、,，;；/]", haystack) if token]
    if any(token in names for token in tokens):
        return []

    has_category = any(re.search(MAJOR_CATEGORY_PATTERN, token) for token in tokens)
    if not has_category and any(name in token for name in names for token in tokens):
        return []

    if segmented and not matched_level:
        whole = normalize(raw)
        if any(code in whole for code in codes) or any(name in whole for name in names):
            return []
        return [issue("专业", "专业按学历分段，缺少考生学历对应段，需人工确认", "专业", raw, "review")]

    if has_category:
        return [issue("专业", "专业要求含大类或相关专业表述，脚本不自动判定", "专业", segment, "review")]

    return [issue("专业", "考生专业未出现在职位专业要求中", "专业", segment)]


def check_political(job, conditions):
    raw = job["political_text"]
    compact = normalize(raw)
    if not compact or "不限" in compact:
        return []

    status = conditions.get("political_status", "群众")
    party = ("中共党员", "中共预备党员")
    league = ("共青团员",)

    if "中共党员" in compact and "共青团员" in compact:
        ok = status in party + league
    elif "中共党员" in compact:
        ok = status in party
    elif "共青团员" in compact:
        ok = status in league + party
    else:
        return [issue("政治面貌", "政治面貌要求无法解析，需人工确认", "政治面貌", raw, "review")]

    if not ok:
        return [issue("政治面貌", "考生政治面貌为%s，职位要求%s" % (status, raw),
                      "政治面貌", raw)]
    return []


def check_grassroots(job, conditions):
    raw = job["grassroots_text"]
    required = parse_years(raw)
    if required is None:
        compact = normalize(raw)
        if not compact or "不限" in compact:
            return []
        return [issue("基层年限", "基层工作年限要求无法解析，需人工确认", "基层工作最低年限", raw, "review")]
    if required <= 0:
        return []
    years = conditions.get("grassroots_years") or 0
    if years < required:
        return [issue("基层年限", "考生基层工作%d年，职位要求%d年" % (years, required),
                      "基层工作最低年限", raw)]
    return []


def service_tags(text):
    compact = normalize(text)
    tags = set()
    for tag, words in SERVICE_KEYWORDS:
        if any(word in compact for word in words):
            tags.add(tag)
    return tags


def check_service_program(job, conditions):
    raw = job["service_text"]
    compact = normalize(raw)
    if not compact or "不限" in compact or "无限制" in compact:
        return []

    user_program = normalize(conditions.get("service_program", "无"))
    if user_program in ("", "无", "否", "没有"):
        return [issue("服务项目", "职位要求服务基层项目经历，考生无此类经历",
                      "服务基层项目工作经历", raw)]

    required_tags = service_tags(compact)
    user_tags = service_tags(user_program)
    if required_tags and user_tags and (required_tags & user_tags):
        return []
    if not required_tags and (user_program in compact or compact in user_program):
        return []
    if not required_tags:
        return [issue("服务项目", "服务基层项目要求无法解析，需人工确认",
                      "服务基层项目工作经历", raw, "review")]
    return [issue("服务项目", "考生服务基层项目经历不符合职位要求",
                  "服务基层项目工作经历", raw)]


def check_region(job, conditions):
    regions = [normalize(region) for region in conditions.get("regions") or [] if region]
    if not regions:
        return []

    haystack = normalize("%s %s" % (job["work_location"], job["settle_location"]))
    if not haystack:
        return [issue("地区", "职位未填写工作地点，无法核对意向地区", "工作地点", job["position_name"], "review")]

    for region in regions:
        if region in haystack:
            return []
        if len(region) >= 2 and haystack in region:
            return []
    return [issue("地区", "工作地点不在考生意向地区", "工作地点",
                  "%s（工作地点）%s（落户地点）" % (job["work_location"], job["settle_location"]))]


def check_fresh(job, conditions):
    remark = job["remark"]
    if not remark:
        return []
    fresh = bool(conditions.get("fresh_graduate"))

    negated = evidence_for(remark, r"非应届")
    if negated:
        if fresh:
            return [issue("应届", "职位要求非应届，考生为应届", "备注", negated)]
        return []

    years = [int(value) for value in re.findall(r"(20\d{2})\s*届", remark)]
    mentions_fresh = "应届" in remark
    if not years and not mentions_fresh:
        return []

    sentence = evidence_for(remark, r"20\d{2}\s*届|应届")
    if years:
        if not fresh:
            return [issue("应届", "职位限定毕业年份，考生非应届", "备注", sentence)]
        graduate_year = conditions.get("graduate_year")
        if graduate_year and int(graduate_year) not in years:
            return [issue("应届", "职位限%d届，考生毕业年份为%s" % (
                years[0], graduate_year), "备注", sentence)]
        if not graduate_year:
            return [issue("应届", "职位限定具体毕业年份，考生未填写毕业年份", "备注", sentence, "review")]
        return []

    if mentions_fresh and not fresh:
        return [issue("应届", "职位要求应届毕业生，考生非应届", "备注", sentence)]
    return []


def check_english(job, conditions):
    remark = job["remark"]
    if not remark:
        return []

    sentence_six = evidence_for(remark, r"(大学英语|英语|CET)[^。；;]{0,20}(六级|CET-?6)")
    sentence_four = evidence_for(remark, r"(大学英语|英语|CET)[^。；;]{0,20}(四级|CET-?4)")
    if not sentence_six and not sentence_four:
        return []

    # 英语专业四级/八级是另一类证书，交给证书复核，不按大学英语四六级硬判。
    if (sentence_six and "专业八级" in sentence_six) or (sentence_four and "专业四级" in sentence_four):
        return []

    english = conditions.get("english") or {}
    cet4 = bool(english.get("cet4"))
    cet6 = bool(english.get("cet6"))

    if sentence_six and sentence_six == sentence_four and "或" in sentence_six:
        if not cet4 and not cet6:
            return [issue("英语", "职位要求英语四级或六级，考生均未通过", "备注", sentence_six)]
        return []
    if sentence_six and not cet6:
        return [issue("英语", "职位要求英语六级，考生未通过", "备注", sentence_six)]
    if sentence_four and not (cet4 or cet6):
        return [issue("英语", "职位要求英语四级，考生未通过", "备注", sentence_four)]
    return []


def check_service_years(job, conditions):
    remark = job["remark"]
    sentence = evidence_for(remark, r"(最低服务年限|最低服务期|服务年限|服务期)")
    if not sentence:
        return []
    if "不设" in sentence or re.search(r"无[^。；;]{0,4}(服务年限|服务期)", sentence):
        return []

    years = parse_years(sentence)
    if years is None:
        return [issue("服务年限", "最低服务年限表述未给出可计算年限，需人工确认", "备注", sentence, "review")]

    if not conditions.get("accept_service_years", True):
        return [issue("服务年限", "职位设最低服务年限%d年，考生不接受" % years, "备注", sentence)]

    max_years = conditions.get("max_service_years")
    if max_years is not None and years > int(max_years):
        return [issue("服务年限", "职位最低服务年限%d年，超过考生可接受的%d年" % (
            years, int(max_years)), "备注", sentence)]
    return []


def check_intensity(job, conditions):
    if conditions.get("accept_intensity", True):
        return []
    remark = job["remark"]
    sentence = evidence_for(remark, INTENSITY_PATTERN)
    if not sentence:
        return []
    return [issue("工作强度", "职位备注含高强度或出差值班表述，考生不接受", "备注", sentence)]


def review_gender(job):
    sentence = evidence_for(job["remark"], GENDER_PATTERN)
    if not sentence:
        return []
    if "男女不限" in sentence or "不限性别" in sentence or "性别不限" in sentence:
        return []
    return [issue("性别表述", "备注出现性别表述，需向招录单位确认适用范围", "备注", sentence)]


def review_hukou(job):
    sentence = evidence_for(job["remark"], HUKOU_PATTERN)
    if not sentence or "不限" in sentence:
        return []
    return [issue("户籍生源", "备注出现户籍或生源要求，需人工核对", "备注", sentence)]


def review_certificate(job):
    sentence = evidence_for(job["remark"], CERT_PATTERN)
    if not sentence:
        return []
    return [issue("证书", "备注要求资格证书，需核对证书类别与取得时间", "备注", sentence)]


def review_physical(job):
    sentence = evidence_for(job["remark"], PHYSICAL_PATTERN)
    if not sentence:
        return []
    return [issue("体检体能", "备注涉及体检或体能要求，需核对标准", "备注", sentence)]


def review_political(job):
    sentence = evidence_for(job["remark"], POLITICAL_PATTERN)
    if not sentence:
        return []
    return [issue("政治考察", "备注涉及政治考察，需确认考察口径", "备注", sentence)]


HARD_CHECKS = [
    check_education,
    check_degree,
    check_major,
    check_political,
    check_grassroots,
    check_service_program,
    check_region,
    check_fresh,
    check_english,
    check_service_years,
    check_intensity,
]

REVIEW_CHECKS = [
    review_gender,
    review_hukou,
    review_certificate,
    review_physical,
    review_political,
]


def dedupe(issues):
    seen = set()
    result = []
    for item in issues:
        key = (item["category"], item["evidence"])
        if key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


def evaluate(job, conditions):
    fails = []
    reviews = []
    for check in HARD_CHECKS:
        for item in check(job, conditions) or []:
            if item.get("level") == "review":
                reviews.append(item)
            else:
                fails.append(item)
    for check in REVIEW_CHECKS:
        for item in check(job) or []:
            reviews.append(item)

    fails = dedupe(fails)
    reviews = dedupe(reviews)
    if fails:
        tier = "不推荐"
    elif reviews:
        tier = "待复核"
    else:
        tier = "优先推荐"

    details = []
    for item in fails:
        details.append("[%s] %s" % (item["category"], item["detail"]))
    for item in reviews:
        details.append("[复核·%s] %s" % (item["category"], item["detail"]))
    if not details:
        details.append("硬条件全部通过，备注无待确认项")

    categories = []
    for item in reviews:
        if item["category"] not in categories:
            categories.append(item["category"])

    evidence = []
    for item in reviews:
        if item["evidence"] and item["evidence"] not in evidence:
            evidence.append(item["evidence"])

    return {
        "tier": tier,
        "fails": fails,
        "reviews": reviews,
        "details": "；".join(details),
        "review_categories": "、".join(categories),
        "review_evidence": "；".join("「%s」" % text for text in evidence),
    }


def load_conditions(path):
    with open(path, "r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise SystemExit("条件文件必须是 JSON 对象")
    conditions = dict(CONDITION_DEFAULTS)
    conditions.update(data)
    if not conditions.get("education"):
        raise SystemExit("条件文件缺少 education（考生最高学历）")
    if user_edu_level(conditions) == 0:
        raise SystemExit("education 取值无法识别，请写 大专/本科/硕士/博士")
    if not conditions.get("major") and not conditions.get("major_codes"):
        sys.stderr.write("提示：未填写 major 或 major_codes，专业相关职位将大量进入待复核\n")
    return conditions


def read_jobs(path):
    jobs, meta = load_jobs(path)
    if not jobs:
        raise SystemExit("没有读到职位数据，请确认文件与表头")
    return jobs, meta


def find_bundled_table():
    """返回 assets 下内置的招考简章；没有则返回 None。"""
    skill_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    assets = os.path.join(skill_root, "assets")
    candidates = sorted(glob.glob(os.path.join(assets, "*.xls")) +
                        glob.glob(os.path.join(assets, "*.xlsx")))
    return candidates[0] if candidates else None


def write_xlsx(rows, summary, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise SystemExit("输出 xlsx 需要 openpyxl，请先执行 pip install openpyxl")

    columns = [
        ("来源表", "source_sheet", 26),
        ("行号", "row_no", 7),
        ("推荐档位", "tier", 10),
        ("部门代码", "dept_code", 10),
        ("部门名称", "dept_name", 28),
        ("用人司局", "bureau", 26),
        ("机构性质", "org_nature", 18),
        ("招考职位", "position_name", 30),
        ("职位属性", "position_attr", 10),
        ("职位分布", "position_dist", 16),
        ("职位简介", "position_desc", 40),
        ("职位代码", "job_code", 15),
        ("机构层级", "level", 12),
        ("考试类别", "exam_type", 12),
        ("招考人数", "headcount", 8),
        ("专业", "major_text", 42),
        ("学历", "education_text", 16),
        ("学位", "degree_text", 14),
        ("政治面貌", "political_text", 14),
        ("基层工作最低年限", "grassroots_text", 14),
        ("服务基层项目工作经历", "service_text", 22),
        ("是否在面试阶段组织专业能力测试", "special_test", 14),
        ("面试人员比例", "interview_ratio", 10),
        ("工作地点", "work_location", 20),
        ("落户地点", "settle_location", 16),
        ("备注", "remark", 60),
        ("部门网站", "dept_site", 26),
        ("咨询电话1", "phone1", 18),
        ("咨询电话2", "phone2", 18),
        ("咨询电话3", "phone3", 18),
        ("判断说明", "details", 60),
        ("复核事项", "review_categories", 22),
        ("复核依据", "review_evidence", 60),
    ]

    book = Workbook()
    sheet = book.active
    sheet.title = "筛选结果"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top")

    for index, (title, _, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row_index, item in enumerate(rows, start=2):
        for column_index, (_, key, _) in enumerate(columns, start=1):
            value = item.get(key, "")
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if key in ("row_no", "headcount"):
                cell.alignment = center
            elif key in ("recommendation", "tier", "review_categories"):
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                cell.alignment = wrap
        tier_cell = sheet.cell(row=row_index, column=3)
        tier_cell.fill = PatternFill("solid", fgColor=TIER_FILL.get(item["tier"], "FFFFFF"))
        tier_cell.font = Font(bold=True, color=TIER_FONT.get(item["tier"], "000000"))

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = "A1:%s%d" % (
        get_column_letter(len(columns)), max(1, len(rows) + 1))

    summary_sheet = book.create_sheet("汇总")
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 18
    summary_sheet.column_dimensions["C"].width = 14
    summary_sheet.column_dimensions["D"].width = 30

    lines = [
        ("国考职位筛选汇总", ""),
        ("生成时间", summary["generated_at"]),
        ("输入文件", summary["input"]),
        ("条件文件", summary["conditions"]),
        ("总职位数", summary["total_jobs"]),
        ("优先推荐", summary["by_tier"].get("优先推荐", 0)),
        ("待复核", summary["by_tier"].get("待复核", 0)),
        ("不推荐", summary["by_tier"].get("不推荐", 0)),
        ("不同职位代码数", summary["unique_job_codes"]),
        ("重复职位代码数", summary["duplicate_job_codes"]),
        ("", ""),
        ("工作表", "表头行", "职位数", "说明"),
    ]
    for index, line in enumerate(lines, start=1):
        for column_index, value in enumerate(line, start=1):
            cell = summary_sheet.cell(row=index, column=column_index, value=value)
            if index == 1:
                cell.font = Font(bold=True, size=14)
            elif index == 12:
                cell.font = header_font
                cell.fill = header_fill
            elif index > 12:
                pass
    row_pointer = 13
    for item in summary["sheets"]:
        summary_sheet.cell(row=row_pointer, column=1, value=item["name"])
        summary_sheet.cell(row=row_pointer, column=2, value=item["header_row"])
        summary_sheet.cell(row=row_pointer, column=3, value=item["jobs"])
        summary_sheet.cell(row=row_pointer, column=4, value="已读取")
        row_pointer += 1
    for name in summary["skipped_sheets"]:
        summary_sheet.cell(row=row_pointer, column=1, value=name)
        summary_sheet.cell(row=row_pointer, column=2, value="-")
        summary_sheet.cell(row=row_pointer, column=3, value=0)
        summary_sheet.cell(row=row_pointer, column=4, value="无表头，已跳过")
        row_pointer += 1

    book.save(path)


def write_review_md(items, summary, path, max_review=0):
    lines = [
        "# 国考职位复核清单",
        "",
        "- 生成时间：%s" % summary["generated_at"],
        "- 输入文件：%s" % summary["input"],
        "- 总职位数：%d；优先推荐 %d；待复核 %d；不推荐 %d" % (
            summary["total_jobs"],
            summary["by_tier"].get("优先推荐", 0),
            summary["by_tier"].get("待复核", 0),
            summary["by_tier"].get("不推荐", 0)),
        "- 本清单只列「待复核」，优先推荐与不推荐见 `筛选结果.xlsx`。",
        "",
    ]

    total = len(items)
    if max_review and total > max_review:
        lines.append("> 待复核共 %d 条，本清单只展示前 %d 条，完整结果见 xlsx。" % (total, max_review))
        lines.append("")
        items = items[:max_review]

    groups = OrderedDict()
    for item in items:
        key = (item["dept_code"], item["dept_name"])
        groups.setdefault(key, []).append(item)

    for (dept_code, dept_name), group in groups.items():
        lines.append("## %s（%s）" % (dept_name or "未命名部门", dept_code or "无代码"))
        lines.append("")
        for item in group:
            lines.append("- **%s｜%s**（职位代码 `%s`，来源表：%s，Excel 行号：%s）" % (
                item["bureau"] or "未知司局",
                item["position_name"] or "未命名职位",
                item["job_code"],
                item["source_sheet"],
                item["row_no"]))
            category_text = item["review_categories"] or "无"
            lines.append("  - 复核事项：%s" % category_text)
            phones = " / ".join(value for value in
                                (item.get("phone1"), item.get("phone2"), item.get("phone3")) if value)
            if phones:
                lines.append("  - 咨询电话：%s" % phones)
            for review in item["reviews"]:
                lines.append("  - %s：%s" % (review["category"], review["detail"]))
            if item["review_evidence"]:
                lines.append("  - 依据：%s" % item["review_evidence"])
        lines.append("")

    with open(path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))


def build_summary(jobs, evaluated, meta, conditions, input_path, conditions_path):
    tier_counter = Counter(item["tier"] for item in evaluated)
    code_counter = Counter(job["job_code"] for job in jobs)
    duplicate_codes = sum(1 for code, count in code_counter.items() if code and count > 1)

    sheet_tiers = Counter()
    for item in evaluated:
        sheet_tiers[(item["source_sheet"], item["tier"])] += 1

    sheets = []
    for sheet in meta["sheets"]:
        sheets.append({
            "name": sheet["name"],
            "header_row": sheet["header_row"],
            "jobs": sheet["jobs"],
            "columns_mapped": sheet["columns_mapped"],
            "tiers": {
                tier: sheet_tiers.get((sheet["name"], tier), 0)
                for tier in ("优先推荐", "待复核", "不推荐")
            },
        })

    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "input": os.path.abspath(input_path),
        "conditions": os.path.abspath(conditions_path) if conditions_path else "",
        "total_jobs": len(jobs),
        "by_tier": {
            tier: tier_counter.get(tier, 0)
            for tier in ("优先推荐", "待复核", "不推荐")
        },
        "unique_job_codes": len([code for code in code_counter if code]),
        "duplicate_job_codes": duplicate_codes,
        "sheets": sheets,
        "skipped_sheets": meta["skipped_sheets"],
        "conditions_snapshot": conditions,
        "rules": [
            "专业大类不自动判定",
            "待复核必须附原文依据",
            "不生成上岸概率、竞争比、分数线预测",
        ],
    }


def find_lookup(jobs, query):
    query = query.strip()
    if "/" in query:
        department, _, job_code = query.partition("/")
        department = code_text(department, 6)
        job_code = code_text(job_code, 12)
        return [job for job in jobs
                if job["dept_code"] == department and job["job_code"] == job_code]
    job_code = code_text(query, 12)
    return [job for job in jobs if job["job_code"] == job_code]


def print_job(job, result=None):
    print("=" * 72)
    print("%s  %s · %s" % (job["job_key"], job["dept_name"], job["bureau"]))
    print("职位：%s" % job["position_name"])
    if result:
        print("档位：%s" % result["tier"])
    print("来源表：%s  行号：%s" % (job["source_sheet"], job["row_no"]))
    print("招考人数：%s  机构层级：%s  考试类别：%s" % (
        job["headcount"], job["level"], job["exam_type"]))
    print("专业：%s" % job["major_text"])
    print("学历：%s  学位：%s  政治面貌：%s" % (
        job["education_text"], job["degree_text"], job["political_text"]))
    print("基层工作最低年限：%s  服务基层项目：%s" % (
        job["grassroots_text"], job["service_text"]))
    print("工作地点：%s  落户地点：%s" % (job["work_location"], job["settle_location"]))
    print("备注：%s" % job["remark"])
    if result:
        print("-" * 72)
        print(result["details"])
        if result["review_evidence"]:
            print("复核依据：%s" % result["review_evidence"])


def main():
    parser = argparse.ArgumentParser(description="国考职位筛选")
    parser.add_argument("--input",
                        help="招考简章 .xls/.xlsx；省略时自动使用 assets 下的内置岗位表")
    parser.add_argument("--conditions", help="考生条件 JSON")
    parser.add_argument("--outdir", default="output", help="输出目录")
    parser.add_argument("--lookup", help="查询 部门代码/职位代码，或只给职位代码")
    parser.add_argument("--max-review", type=int, default=0,
                        help="复核清单最多输出条数，0 表示全部")
    args = parser.parse_args()

    input_path = args.input
    if not input_path:
        input_path = find_bundled_table()
        if not input_path:
            parser.error("未找到内置岗位表，请用 --input 指定招考简章")
        print("使用内置岗位表：%s" % os.path.abspath(input_path))

    jobs, meta = read_jobs(input_path)

    if args.lookup:
        matches = find_lookup(jobs, args.lookup)
        if not matches:
            print("未找到：%s" % args.lookup)
            return 1
        conditions = load_conditions(args.conditions) if args.conditions else None
        for job in matches:
            result = evaluate(job, conditions) if conditions else None
            print_job(job, result)
        return 0

    if not args.conditions:
        parser.error("筛选模式需要 --conditions；只查单条职位请用 --lookup")

    conditions = load_conditions(args.conditions)
    evaluated = []
    for job in jobs:
        result = evaluate(job, conditions)
        item = dict(job)
        item.update({
            "tier": result["tier"],
            "details": result["details"],
            "review_categories": result["review_categories"],
            "review_evidence": result["review_evidence"],
            "reviews": result["reviews"],
        })
        evaluated.append(item)

    evaluated.sort(key=lambda item: (
        TIER_ORDER.get(item["tier"], 9),
        item["dept_code"],
        item["job_code"],
    ))

    summary = build_summary(jobs, evaluated, meta, conditions, input_path, args.conditions)
    os.makedirs(args.outdir, exist_ok=True)
    xlsx_path = os.path.join(args.outdir, "筛选结果.xlsx")
    md_path = os.path.join(args.outdir, "复核清单.md")
    json_path = os.path.join(args.outdir, "run_summary.json")

    write_xlsx(evaluated, summary, xlsx_path)
    review_items = [item for item in evaluated if item["tier"] == "待复核"]
    write_review_md(review_items, summary, md_path, args.max_review)
    with open(json_path, "w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=2)

    print("总职位数 %d" % summary["total_jobs"])
    for tier in ("优先推荐", "待复核", "不推荐"):
        print("%s %d" % (tier, summary["by_tier"].get(tier, 0)))
    print("工作表：" + "；".join(
        "%s %d 条（表头第 %d 行）" % (item["name"], item["jobs"], item["header_row"])
        for item in summary["sheets"]))
    if summary["skipped_sheets"]:
        print("跳过空表：" + "、".join(summary["skipped_sheets"]))
    print("输出：%s" % os.path.abspath(args.outdir))
    return 0


if __name__ == "__main__":
    sys.exit(main())
